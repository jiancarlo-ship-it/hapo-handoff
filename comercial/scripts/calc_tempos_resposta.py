"""
Calcula os 3 tempos de resposta pedidos pelo Jian:
  1) Criado em -> Primeira vez que entrou na fase Contato Inicial   (tempo de atendimento)
  2) Contato Inicial -> Reuniao Agendada                             (tempo até agendar)
  3) Reuniao Comercial (fase, funil de vendas) -> Negocio Fechado    (tempo até fechar após reunião)

Usa os exports já baixados hoje (13/08/2026). "Todas as unidades" é sempre a soma dos
valores brutos das 3 unidades nomeadas (Goalfy/Educação/Assessoria), não uma query
separada sobre a base toda — isso evita puxar histórico de meses muito antigos ou
leads/negócios sem unidade classificada, que inflam ou distorcem os números.
"""
import json
import os
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(__file__)
RAIZ = os.path.join(AQUI, '..')


def parse_dt(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, str):
        for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y'):
            try:
                return datetime.strptime(v, fmt)
            except ValueError:
                pass
    return None


def normaliza_unidade(v):
    if not v:
        return None
    vl = v.strip().lower()
    if 'goalfy' in vl or 'goslfy' in vl:
        return 'Goalfy'
    if 'assessoria' in vl:
        return 'Hapo Assessoria'
    if 'educa' in vl:
        return 'Hapo Educação'
    return None


def classificar_unidade_prevendas(g):
    v = normaliza_unidade(g('Unidade de Negócio'))
    if v:
        return v
    v = normaliza_unidade(g('Etiqueta - Unidade de Negócio'))
    if v:
        return v
    for campo in ('Campanha', 'Fonte do Lead', 'Título do Evento'):
        val = g(campo)
        if not val:
            continue
        vl = val.lower()
        if 'goalfy' in vl:
            return 'Goalfy'
        if 'assessoria' in vl:
            return 'Hapo Assessoria'
        if 'educa' in vl or 'imers' in vl:
            return 'Hapo Educação'
    return 'Não classificado'


def stats(vals):
    vals = [v for v in vals if v is not None and v >= 0]
    if not vals:
        return None
    vals.sort()
    n = len(vals)
    return {'n': n, 'mediana': round(vals[n // 2], 2), 'media': round(sum(vals) / n, 2)}


def carregar(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb['Sheet']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]
    return rows, idx


def no_mes(v, ano, mes):
    d = parse_dt(v)
    return d is not None and d.year == ano and d.month == mes


rows_pv, idx_pv = carregar(os.path.join(RAIZ, 'Exports (auto)/funil_2026-08-13.xlsx'))
def g_pv(r, h):
    return r[idx_pv[h]]


is_base_antiga = lambda r: (g_pv(r, 'Fonte do Lead') or '').__contains__('Base Antiga')
ativo = [r for r in rows_pv if not is_base_antiga(r)]

rows_ve, idx_ve = carregar(os.path.join(RAIZ, 'Exports (auto)/funil_vendas_2026-08-13.xlsx'))
def g_ve(r, h):
    return r[idx_ve[h]]


MESES = [('2026-07', 2026, 7), ('2026-08', 2026, 8)]
UNIDADES = ['Goalfy', 'Hapo Educação', 'Hapo Assessoria']

# vals_por_unidade[metrica][unidade][mes] = lista de dias (bruta)
vals = {m: {u: {p: [] for p, _, _ in MESES} for u in UNIDADES} for m in ('atendimento', 'agendamento', 'fechamento')}

for r in ativo:
    u = classificar_unidade_prevendas(lambda h, r=r: g_pv(r, h))
    if u not in UNIDADES:
        continue
    criado = g_pv(r, 'Criado em')
    c = parse_dt(criado)
    ci = parse_dt(g_pv(r, 'Primeira vez que entrou na fase Contato Inicial'))
    ra = parse_dt(g_pv(r, 'Primeira vez que entrou na fase Reunião Agendada'))
    for pkey, ano, mes in MESES:
        if not no_mes(criado, ano, mes):
            continue
        if c and ci and ci >= c:
            vals['atendimento'][u][pkey].append((ci - c).total_seconds() / 86400)
        if ci and ra and ra >= ci:
            vals['agendamento'][u][pkey].append((ra - ci).total_seconds() / 86400)

for r in rows_ve:
    u = normaliza_unidade(g_ve(r, 'Unidade de Negócio'))
    if u not in UNIDADES:
        continue
    criado = g_ve(r, 'Criado em')
    rc = parse_dt(g_ve(r, 'Primeira vez que entrou na fase Reunião Comercial'))
    nf = parse_dt(g_ve(r, 'Primeira vez que entrou na fase Negócio Fechado'))
    for pkey, ano, mes in MESES:
        if not no_mes(criado, ano, mes):
            continue
        if rc and nf and nf >= rc:
            vals['fechamento'][u][pkey].append((nf - rc).total_seconds() / 86400)

resultado = {}
for metrica in vals:
    resultado[metrica] = {}
    combinado_todas = {p: [] for p, _, _ in MESES}
    for u in UNIDADES:
        resultado[metrica][u] = {}
        combinado_u = []
        for pkey, _, _ in MESES:
            resultado[metrica][u][pkey] = stats(vals[metrica][u][pkey])
            combinado_todas[pkey].extend(vals[metrica][u][pkey])
            combinado_u.extend(vals[metrica][u][pkey])
        resultado[metrica][u]['jul_ago'] = stats(combinado_u)
    resultado[metrica]['Todas'] = {}
    combinado_todas_geral = []
    for pkey, _, _ in MESES:
        resultado[metrica]['Todas'][pkey] = stats(combinado_todas[pkey])
        combinado_todas_geral.extend(combinado_todas[pkey])
    resultado[metrica]['Todas']['jul_ago'] = stats(combinado_todas_geral)

out_path = os.path.join(RAIZ, 'Analises/2026-08-13/tempos_resposta.json')
with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(resultado, f, ensure_ascii=False, indent=1)
print('OK ->', out_path)
