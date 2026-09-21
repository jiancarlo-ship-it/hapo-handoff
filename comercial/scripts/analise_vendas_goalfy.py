"""
Analise do funil de VENDAS (pos pre-vendas) a partir do export Excel do Goalfy.

Pipeline: Oportunidade (SQL) -> Reuniao Comercial -> Em Contato -> Proposta
Comercial -> Follow-Up -> Negociacao/Contrato -> Negocio Fechado / Negocio
Perdido -> Base Geral (pool de descanso).

Uso:
    python analise_vendas_goalfy.py caminho.xlsx --mes 2026-07 --unidade Goalfy --out saida.json

Sem --mes/--unidade, roda sobre a base inteira.
"""
import argparse
import json
from collections import Counter
from datetime import datetime

import openpyxl


def parse_dt(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, str):
        for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y'):
            try:
                return datetime.strptime(v, fmt)
            except ValueError:
                pass
    return None


def to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace('R$', '').replace('.', '').replace(',', '.').strip()
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _normaliza_unidade(v):
    if not v:
        return None
    vl = v.strip().lower()
    if 'goalfy' in vl or 'goslfy' in vl:
        return 'Goalfy'
    if 'assessoria' in vl:
        return 'Hapo Assessoria'
    if 'educa' in vl:
        return 'Hapo Educação'
    return None


def carregar_linhas(caminho_xlsx):
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb['Sheet']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]
    return rows, idx


def filtrar_mes(rows, idx, ano_mes):
    ano, mes = (int(x) for x in ano_mes.split('-'))
    c = idx['Criado em']

    def dt_no_mes(v):
        d = parse_dt(v)
        return d is not None and d.year == ano and d.month == mes

    return [r for r in rows if dt_no_mes(r[c])]


def filtrar_unidade(rows, idx, unidade):
    c = idx['Unidade de Negócio']
    return [r for r in rows if _normaliza_unidade(r[c]) == unidade]


def analisar(rows, idx):
    def g(r, h):
        return r[idx[h]]

    out = {'total_linhas': len(rows)}
    out['fase_atual'] = Counter(g(r, 'Fase Atual') for r in rows).most_common(15)
    out['unidade_negocio_bruta'] = Counter(g(r, 'Unidade de Negócio') for r in rows).most_common(10)

    # IMPORTANTE: usar "ever entered" (Primeira vez que entrou na fase X), nao
    # Fase Atual — 42% dos cards acabam arquivados em "Base Geral" depois de
    # fechar/perder, entao Fase Atual subestima MUITO fechado/perdido.
    fechado = [r for r in rows if g(r, 'Primeira vez que entrou na fase Negócio Fechado') is not None]
    perdido = [r for r in rows if g(r, 'Primeira vez que entrou na fase Negócio Perdido') is not None]
    decididos = len(fechado) + len(perdido)
    out['negocio_fechado'] = len(fechado)
    out['negocio_perdido'] = len(perdido)
    out['taxa_fechamento_pct'] = round(100 * len(fechado) / decididos, 1) if decididos else None

    # comparativo com a leitura ingenua por Fase Atual, so pra documentar o quanto
    # Base Geral escondia
    out['_fase_atual_fechado'] = sum(1 for r in rows if g(r, 'Fase Atual') == 'Negócio Fechado')
    out['_fase_atual_perdido'] = sum(1 for r in rows if g(r, 'Fase Atual') == 'Negócio Perdido')

    out['em_andamento_sem_decisao'] = len(rows) - decididos

    for fase in ('Oportunidade (SQL)', 'Reunião Comercial', 'Em Contato', 'Proposta Comercial', 'Follow-Up', 'Negociação/Contrato'):
        col = f'Primeira vez que entrou na fase {fase}'
        out.setdefault('ever_entered', {})[fase] = sum(1 for r in rows if g(r, col) is not None)

    valores = [to_float(g(r, 'Valor da Proposta')) for r in fechado]
    valores = [v for v in valores if v]
    out['ticket_medio_fechado'] = round(sum(valores) / len(valores), 2) if valores else None
    out['ticket_n'] = len(valores)
    out['receita_total_fechado'] = round(sum(valores), 2) if valores else None

    out['motivo_nao_fechamento'] = Counter(
        (g(r, 'Motivo de Não Fechamento') or '').strip() for r in perdido if g(r, 'Motivo de Não Fechamento')
    ).most_common(15)

    # lead time: criado -> negocio fechado
    lt = []
    for r in fechado:
        c = parse_dt(g(r, 'Criado em'))
        f = parse_dt(g(r, 'Primeira vez que entrou na fase Negócio Fechado'))
        if c and f:
            lt.append((f - c).total_seconds() / 86400)
    if lt:
        lt.sort()
        n = len(lt)
        out['ciclo_venda_dias'] = {'n': n, 'mediana': round(lt[n // 2], 2), 'media': round(sum(lt) / n, 2)}

    return out


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('xlsx')
    parser.add_argument('--mes', help='AAAA-MM')
    parser.add_argument('--unidade', help='Goalfy | Hapo Educação | Hapo Assessoria')
    parser.add_argument('--out', default='analise_vendas_saida.json')
    args = parser.parse_args()

    rows, idx = carregar_linhas(args.xlsx)

    if args.mes:
        rows = filtrar_mes(rows, idx, args.mes)
    if args.unidade:
        rows = filtrar_unidade(rows, idx, args.unidade)

    resultado = analisar(rows, idx)
    with open(args.out, 'w', encoding='utf-8') as f:
        json.dump(resultado, f, ensure_ascii=False, indent=1, default=str)
    print(f'OK - {len(rows)} linhas analisadas -> {args.out}')


if __name__ == '__main__':
    main()
